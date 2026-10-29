#!/usr/bin/env python3
"""Supplementary Table S18 — gene-level annotation behind the temperature toolkits.

Supplementary Fig. S27 groups the strain-specific orthogroups by their HOG-level
Pfam annotation. A HOG annotation is however pooled over all 21 genomes, so it
does not say which gene a given strain carries or what that gene encodes. This
table supplies that: every gene behind the toolkit orthogroups, resolved per
strain to its own dbCAN, Pfam and eggNOG annotation.

Sheets
  per_strain_gene  one row per gene, for every strain that carries the
                   orthogroup: gene id, eggNOG name, COG class, Pfam domains,
                   CAZy class and KEGG KO of that gene in that strain
  HOG_summary      one row per orthogroup: the HOG-level union annotation
                   alongside what is actually observed at gene level, and how
                   many of its genes carry no Pfam record at all
  definitions      how the toolkits were defined and how to read the two levels

Input : 05.pan-genome/orthofinder-shorter/N0.gene_vs_HOG.tsv
        05.pan-genome/orthofinder-shorter/N0.HOG.hog_pav.tsv
        05.pan-genome/orthofinder-shorter/N0.HOG_annotation.v2.tsv
        02.annotation/per_isolates/<iso>/<iso>.gene_{pfam,dbcan,eggnog}.tsv
        bmu_ref_groups.211.txt
Output: 00.manuscript/supplementary/suppl_table.S18_toolkit_gene_annotation.xlsx
"""
import os

import pandas as pd
from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(ROOT, "00.manuscript/supplementary",
                   "suppl_table.S18_toolkit_gene_annotation.xlsx")

STRAINS = ["BMU" + x.zfill(5) for x in
           ["9608", "7029", "7936", "1336", "2406", "5572", "2479", "2031", "12685",
            "12686", "1287", "1935", "1936", "12683", "12684", "2407", "3976", "2028",
            "1973", "2745", "10147"]]
# groups come from the per-strain CMH test (P < 0.05, stratified by inoculum dose),
# not from a hard cutoff on the continuous difference in mean mortality
_bias = pd.read_csv(os.path.join(ROOT, "07.lethal_groups/phenotype_analysis/"
                                       "temperature_bias_cmh.tsv"), sep="\t")
B37 = sorted(_bias.strain[_bias["class"] == "37C_biased"])
B28 = sorted(_bias.strain[_bias["class"] == "28C_biased"])
GROUP = {s: "37C-biased" for s in B37}
GROUP.update({s: "28C-biased" for s in B28})
GROUP.update({s: "neutral" for s in STRAINS if s not in GROUP})
ANNOT = os.path.join(ROOT, "05.pan-genome/orthofinder-shorter/N0.HOG_annotation.v2.tsv")
PER = os.path.join(ROOT, "02.annotation/per_isolates")

pav = pd.read_csv(os.path.join(ROOT, "05.pan-genome/orthofinder-shorter/"
                                     "N0.HOG.hog_pav.tsv"), sep="\t", dtype={"Present": "int8"})
wide = (pav[pav.SampleID.isin(STRAINS)]
        .pivot(index="SampleID", columns="HOG_ID", values="Present").fillna(0).loc[STRAINS])
prev = wide.sum(0)
keep = wide.loc[:, (prev >= 3) & (prev <= 19)]
mid = [s for s in STRAINS if s not in B37 + B28]
sp37 = list(keep.columns[(keep.loc[B37].sum(0) == len(B37)) &
                         (keep.loc[B28 + mid].sum(0) <= 3)])
sp28 = list(keep.columns[(keep.loc[B28].sum(0) == len(B28)) &
                         (keep.loc[B37 + mid].sum(0) <= 4)])
TOOLKIT = {h: "37C-biased" for h in sp37}
TOOLKIT.update({h: "28C-biased" for h in sp28})
print(f"toolkit: 37C-biased {len(sp37)} HOGs · 28C-biased {len(sp28)} HOGs")

ann = pd.read_csv(ANNOT, sep="\t").drop_duplicates("HOG_ID").set_index("HOG_ID")
gvh = pd.read_csv(os.path.join(ROOT, "05.pan-genome/orthofinder-shorter/"
                                     "N0.gene_vs_HOG.tsv"),
                  sep="\t", usecols=["GeneID", "SampleID", "HOG_ID"])
gvh = gvh[gvh.HOG_ID.isin(TOOLKIT) & gvh.SampleID.isin(STRAINS)]

src = pd.read_csv(os.path.join(ROOT, "bmu_ref_groups.211.txt"), sep="\t")
src_map = dict(zip(src.BMU_ID.astype(str), src.MedEnv.astype(str)))


def load(iso, kind):
    if kind == "cazy":
        d = pd.read_csv(f"{PER}/{iso}/{iso}.gene_dbcan.tsv", sep="\t", header=None,
                        names=["g", "v"])
        return d.groupby("g")["v"].apply(lambda s: ";".join(sorted(set(s)))).to_dict()
    if kind == "pfam":
        d = pd.read_csv(f"{PER}/{iso}/{iso}.gene_pfam.tsv", sep="\t", header=None,
                        usecols=[0, 2], names=["g", "dom"])
        return d.groupby("g")["dom"].apply(lambda s: ";".join(sorted(set(s)))).to_dict()
    d = pd.read_csv(f"{PER}/{iso}/{iso}.gene_eggnog.tsv", sep="\t", header=None,
                    usecols=[0, 1, 2, 4], names=["g", "name", "cog", "ko"])
    return {r.g: (r.name, r.cog, r.ko) for r in d.itertuples(index=False)}


ann_maps = {i: {k: load(i, k) for k in ("pfam", "cazy", "eggnog")} for i in STRAINS}

rows = []
for r in gvh.itertuples(index=False):
    iso = r.SampleID
    pf = ann_maps[iso]["pfam"].get(r.GeneID)
    cz = ann_maps[iso]["cazy"].get(r.GeneID)
    eg = ann_maps[iso]["eggnog"].get(r.GeneID)
    name, cog, ko = (eg if eg else ("", "", ""))
    rows.append({
        "HOG_ID": r.HOG_ID,
        "toolkit": TOOLKIT[r.HOG_ID],
        "strain": iso,
        "source": src_map.get(iso, ""),
        "strain_temperature_group": GROUP[iso],
        "gene_id": r.GeneID,
        "gene_name_eggnog": name if name not in ("-", "") else "",
        "COG_class": cog if cog not in ("-", "") else "",
        "Pfam_domains_gene": pf or "",
        "CAZy_gene": cz or "",
        "KEGG_KO_gene": ko if ko not in ("-", "") else "",
    })
per = pd.DataFrame(rows).sort_values(["toolkit", "HOG_ID", "strain"])

summary = []
for h, grp in per.groupby("HOG_ID"):
    a = ann.loc[h] if h in ann.index else None
    allg = grp["Pfam_domains_gene"]
    n_no_pfam = int((allg == "").sum())
    doms = sorted({d for s in allg for d in str(s).split(";") if d})
    caz = sorted({c for s in grp["CAZy_gene"] for c in str(s).split(";") if c})
    summary.append({
        "HOG_ID": h,
        "toolkit": TOOLKIT[h],
        "n_strains_carrying": grp["strain"].nunique(),
        "n_genes": len(grp),
        "HOG_level_gene_names": str(a["Gene_Name"]) if a is not None else "",
        "HOG_level_CAZy": str(a["CAZy_Class"]) if a is not None else "",
        "HOG_level_Pfam": str(a["Pfam_Domain"]) if a is not None else "",
        "gene_level_Pfam_observed": ";".join(doms),
        "gene_level_CAZy_observed": ";".join(caz),
        "n_genes_without_Pfam_record": n_no_pfam,
    })
summ = pd.DataFrame(summary).sort_values(["toolkit", "HOG_ID"])

defs = [
    ("item", "value"),
    ("what this table resolves",
     "Supplementary Fig. S27 colours these orthogroups by their HOG-level Pfam annotation. "
     "A HOG annotation is pooled over all 21 genomes, so it does not describe the gene a "
     f"particular strain carries. This table lists, for every strain carrying one of the {len(sp37) + len(sp28)} "
     "toolkit orthogroups, the gene itself and that gene's own annotations."),
    ("toolkit definition",
     f"Strain-specific orthogroups of the temperature-biased strains. 37C-biased = {', '.join(B37)}; "
     f"28C-biased = {', '.join(B28)}. The biased groups are the isolates whose temperature preference is "
     "significant by a Cochran-Mantel-Haenszel test stratified by inoculum dose at P < 0.05, the direction "
     "of the common odds ratio setting the group. "
     f"An orthogroup is 37C-specific when carried by all {len(B37)} 37C-biased strains and by at most 3 of "
     f"the other {21 - len(B37)}; 28C-specific when carried by all {len(B28)} 28C-biased strains and by at "
     f"most 4 of the other {21 - len(B28)}."),
    ("gene assignment",
     "Gene-to-orthogroup assignment from N0.gene_vs_HOG.tsv; the id it lists is a transcript id "
     "(-T1 suffix). One row per gene, so an orthogroup carried as two paralogues in one strain "
     "gives two rows."),
    ("annotation sources",
     "Per isolate: dbCAN (CAZy), Pfam (HMMER domains) and eggNOG (gene name, COG class, KEGG KO), "
     "each read from 02.annotation/per_isolates/<strain>/."),
    ("coverage caveat",
     "Pfam records are available for roughly 70% of genes, so an empty Pfam_domains_gene cell "
     "means no annotation, not an absence of domains. dbCAN is run on all genes and returns only "
     "hits (about 690 per genome, the expected CAZyme complement), so an empty CAZy_gene cell "
     "is a valid negative."),
    ("how the two levels differ",
     "HOG_level_* columns are the union over all members of the orthogroup; gene_level_* columns "
     "are what is actually observed in the strains listed. Where they disagree, the HOG-level "
     "row does not describe the genes of the biased strains."),
]

wb = Workbook()
ws = wb.active
ws.title = "per_strain_gene"
ws.append(list(per.columns))
for row in per.itertuples(index=False):
    ws.append([v.item() if hasattr(v, "item") else v for v in row])

ws2 = wb.create_sheet("HOG_summary")
ws2.append(list(summ.columns))
for row in summ.itertuples(index=False):
    ws2.append([v.item() if hasattr(v, "item") else v for v in row])

ws3 = wb.create_sheet("definitions")
for r in defs:
    ws3.append(list(r))

wb.save(OUT)
print(f"wrote {OUT}")
print(f"  per_strain_gene: {len(per)} 行 · HOG_summary: {len(summ)} 行")
print("\n每个 HOG 的基因数与无 Pfam 记录数：")
print(summ[["HOG_ID", "toolkit", "n_strains_carrying", "n_genes",
            "n_genes_without_Pfam_record", "gene_level_CAZy_observed"]].to_string(index=False))
