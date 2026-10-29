#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Functional enrichment of BMU02483's clinical-only genes, split by the route that supplied them.

Question
    If a transformation happened, which functions came in -- and are the accessory-chromosome
    and core-chromosome routes functionally the same?

Sets (from Fig. 4k / Supplementary Table S14, sheet 'Clinical_only_gene_routes')
    clinical-only  the 1,138 BMU02483 genes whose donors (>=99% identity over 211 isolates)
                   include clinical isolates and none of the 17 nearby environmental isolates
    AC only        229 of them, donors exclusively on accessory chromosomes
    CC only        164 of them, donors exclusively on core chromosomes

Background = all annotated BMU02483 genes (16,864).

Categories
    Pfam primary class (scripts/stats_020_genes.classify_pfam), protein-kinase status,
    CAZyme major class (dbCAN), KEGG BRITE level 2, effector status (EffectorP master).

KEGG NOTE
    scripts/stats_transfer_function.py maps KEGG with `load_kegg_level2(...).get(ko)`, but that
    table is keyed by PATHWAY id, not by KO, so no gene ever received a KEGG category (the
    published column is 100% "No KEGG"). Here the KO -> pathway step is done explicitly, from
    the eggNOG KEGG_Pathway column, so the KEGG categories are actually populated.

Output: sheet 'Clinical_only_gene_functions' in suppl_table.BMU02483_gene_content.xlsx
"""

import collections
import csv
import os
import sys

import openpyxl
import pandas as pd
from scipy import stats as st

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "scripts"))
import stats_020_genes as s020                      # noqa: E402

ISO = "BMU02483"
ANN = f"02.annotation/per_isolates/{ISO}"
KEGG = "08.pan_020_genes/kegg_pathway_categories.tsv"
EFF = "08.pan_020_genes/effector_master.gene_level.tsv"
XLSX = "00.manuscript/supplementary/suppl_table.BMU02483_gene_content.xlsx"
SHEET = "Clinical_only_gene_functions"


def bh(ps):
    n = len(ps)
    order = sorted(range(n), key=lambda i: ps[i])
    q = [0.0] * n
    prev = 1.0
    for rank, i in enumerate(reversed(order), start=1):
        val = min(prev, ps[i] * n / (n - rank + 1))
        q[i] = val
        prev = val
    return q


def main():
    # ---- compartment of every clinical-only gene, read back from the table itself ----
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    ws = wb["Clinical_only_gene_routes"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    ci = {h: i for i, h in enumerate(hdr)}
    grp_of = {}
    for r in rows[1:]:
        grp_of[r[ci["BMU02483_gene"]]] = r[ci["compartment"]]
    wb.close()
    AC = {g for g, c in grp_of.items() if c == "AC only"}
    CC = {g for g, c in grp_of.items() if c == "CC only"}
    CLIN = set(grp_of)
    print(f"clinical-only {len(CLIN)}   AC only {len(AC)}   CC only {len(CC)}")

    # ---- annotations ----
    pfam = collections.defaultdict(list)
    for line in open(f"{ANN}/{ISO}.gene_pfam.tsv"):
        p = line.rstrip("\n").split("\t")
        if len(p) >= 3:
            pfam[p[0]].append(p[2])
    cazy = collections.defaultdict(set)
    for line in open(f"{ANN}/{ISO}.gene_dbcan.tsv"):
        p = line.rstrip("\n").split("\t")
        if len(p) >= 2 and len(p[1]) >= 2:
            cazy[p[0]].add(p[1][:2])
    kegg2 = {}
    for line in open(KEGG):
        p = line.rstrip("\n").split("\t")
        if len(p) >= 3 and p[0] != "pathway_id":
            kegg2[p[0]] = p[2].strip()
    kegg_of = collections.defaultdict(set)
    for line in open(f"{ANN}/{ISO}.gene_eggnog.tsv"):
        p = line.rstrip("\n").split("\t")
        if len(p) < 6 or not p[5].startswith("ko"):
            continue
        for pw in p[5].split(","):
            pw = pw.strip()
            if pw.startswith(("ko", "map")):
                pw = pw[2:]
            if pw in kegg2:
                kegg_of[p[0]].add(kegg2[pw])
    eff = pd.read_csv(EFF, sep="\t")
    eff = eff[eff["SampleID"] == ISO]
    eff_of = dict(zip(eff["GeneID"], eff["IsEffector"].astype(str)))

    # ---- per-gene category sets ----
    universe = pd.read_csv(f"{ANN}/{ISO}.gene_eggnog.tsv", sep="\t", header=None, usecols=[0])[0].unique()
    universe = set(universe)
    print(f"annotated genes: {len(universe)}   genes with KEGG level2: {len(kegg_of)}"
          f"   with effector call: {len(eff_of)}")

    def cats(g):
        out = []
        doms = pfam.get(g, [])
        out.append(("Pfam class", s020.classify_pfam(doms)))
        if any(d.lower().startswith(("pkinase", "hiska")) or "kinase" in d.lower() for d in doms):
            out.append(("Protein kinase", "kinase"))
        if g in cazy:
            for c in cazy[g]:
                out.append(("CAZyme class", c))
        for k in kegg_of.get(g, ()):
            out.append(("KEGG level2", k))
        if eff_of.get(g) in ("True", "true"):
            out.append(("Effector", "effector"))
        return out

    cat_members = collections.defaultdict(set)       # (type, cat) -> genes
    for g in universe:
        for t, c in cats(g):
            cat_members[(t, c)].add(g)

    out_rows = []
    for (ftype, cat), members in sorted(cat_members.items()):
        n1, n2 = len(CLIN), len(universe)
        k1 = len(members & CLIN)
        if k1 < 3:
            continue
        k2 = len(members)
        orr, p = st.fisher_exact([[k1, n1 - k1], [k2 - k1, (n2 - n1) - (k2 - k1)]])
        a1 = len(members & AC)
        c1 = len(members & CC)
        if a1 + c1:
            orr2, p2 = st.fisher_exact([[a1, len(AC) - a1], [c1, len(CC) - c1]])
        else:
            orr2, p2 = float("nan"), float("nan")
        out_rows.append({
            "Functional_type": ftype, "Category": cat,
            "N_clinical_only_with": k1, "Pct_clinical_only": 100 * k1 / n1,
            "OR_vs_genome": orr, "P": p,
            "N_AC_with": a1, "Pct_AC": 100 * a1 / len(AC),
            "N_CC_with": c1, "Pct_CC": 100 * c1 / len(CC),
            "OR_AC_vs_CC": orr2, "P_AC_vs_CC": p2,
        })
    df = pd.DataFrame(out_rows)
    df["q"] = bh(list(df["P"].clip(lower=1e-300)))
    df = df.sort_values(["Functional_type", "P"])

    wb = openpyxl.load_workbook(XLSX)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)
    ws.append(list(df.columns))
    for _, r in df.iterrows():
        ws.append([r[c] if not isinstance(r[c], float) else round(float(r[c]), 6)
                   for c in df.columns])
    wb.save(XLSX)
    print(f"\n{len(df)} categories written to [{SHEET}]")
    top = df[df["q"] < 0.05].sort_values("P").head(12)
    print("\ntop enriched categories in clinical-only genes:")
    print(top[["Functional_type", "Category", "N_clinical_only_with", "Pct_clinical_only",
               "OR_vs_genome", "q", "Pct_AC", "Pct_CC", "P_AC_vs_CC"]].to_string(index=False))


if __name__ == "__main__":
    main()
