#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Are clinically relevant functions enriched in BMU02483 relative to its own clade, and do
they sit on accessory chromosomes?

(a) CARRIAGE: for each Pfam-defined category, compare the proportion of BMU02483's genes
    carrying it with the same proportion in the 17 environmental isolates closest to it in
    the core-genome tree order. Two-sided Fisher exact test on the pooled counts, plus the
    median and range of the 17 individual proportions so that BMU02483 can be seen against
    the spread rather than only against a pooled number.

(b) AC ASSOCIATION: for BMU02483's genes in that category, the compartment of their donors
    at >=99% identity (accessory chromosome / core chromosome / unresolvable, the same
    classification as Fig. 4k), tested against the genome-wide AC rate.

"Drug resistance" is approximated by the efflux pump families (MFS + ABC), because no Pfam
domain is specific to resistance; that is stated wherever the number is used.
"""

import collections
import csv
import os
import statistics

from scipy import stats as st

TGS = "06.accessory-chrs/tgs_assemblies"
ANN = "02.annotation/per_isolates"
FOCUS = "BMU02483"
ISO13 = ["BMU01933", "BMU02028", "BMU02474", "BMU03267", "BMU06673",
         "BMU12683", "BMU12684", "BMU12685", "BMU12686",
         "Fu3.1", "Fu6.1", "SB1", "SX13"]

CATS = {
    "Cytochrome P450": lambda d: any("p450" in x.lower() for x in d),
    "MFS transporter": lambda d: any(x.lower().startswith("mfs") for x in d),
    "Sugar transporter (Sugar_tr)": lambda d: any(x.lower().startswith("sugar_tr") for x in d),
    "ABC transporter": lambda d: any(x.lower().startswith("abc") for x in d),
    "Efflux pumps (MFS + ABC)": lambda d: any(x.lower().startswith(("mfs", "abc")) for x in d),
    "HET domain": lambda d: any(x.lower().startswith("het") for x in d),
    "CFEM domain": lambda d: any("cfem" in x.lower() for x in d),
    "RTA1 domain": lambda d: any("rta1" in x.lower() for x in d),
}


def pfam_of(iso):
    d = collections.defaultdict(list)
    p = f"{ANN}/{iso}/{iso}.gene_pfam.tsv"
    if not os.path.exists(p):
        return d
    for line in open(p):
        f = line.rstrip("\n").split("\t")
        if len(f) >= 3:
            d[f[0]].append(f[2])
    return d


def n_genes(iso):
    n = 0
    for line in open(f"{ANN}/{iso}/funannotate/predict_results/fusarium.gff3"):
        if line.startswith("#"):
            continue
        f = line.rstrip("\n").split("\t")
        if len(f) >= 9 and f[2] == "mRNA":
            n += 1
    return n


def main():
    grp = {r["BMU_ID"]: r for r in csv.DictReader(
        open("05.pan-genome/bmu_ref_groups.211.txt"), delimiter="\t")}
    order = [l.strip() for l in open(f"{TGS}/tree_order_allSC.txt") if l.strip()]
    b = order.index(FOCUS)
    near = order[max(0, b - 10):b + 11]
    NEAR = [s for s in near if grp.get(s, {}).get("Species") == "falciformis"
            and grp[s]["MedEnv"] == "env"]

    # ---- per-isolate Pfam and gene counts -----------------------------------------
    foc_pfam = pfam_of(FOCUS)
    foc_n = n_genes(FOCUS)
    env = []
    for iso in NEAR:
        env.append((iso, pfam_of(iso), n_genes(iso)))
    print(f"BMU02483: {foc_n} genes;  nearby env isolates: {len(env)}")

    # ---- donor compartment of every BMU02483 gene (the Fig. 4k classification) -----
    tid2c = {}
    for iso in ISO13:
        acc = set(open(f"{TGS}/{iso}.accessory_chrs.id").read().split())
        for line in open(f"{ANN}/{iso}/funannotate/predict_results/fusarium.gff3"):
            if line.startswith("#"):
                continue
            f = line.rstrip("\n").split("\t")
            if len(f) < 9 or f[2] != "mRNA":
                continue
            for at in f[8].split(";"):
                if at.startswith("ID="):
                    tid2c[at[3:]] = "AC" if f[0] in acc else "CC"
                    break
    comp = {}
    clusters = collections.defaultdict(list)
    for line in open("tmp/mmseqs99/out99_cluster.tsv"):
        p = line.rstrip("\n").split("\t")
        if len(p) > 1:
            clusters[p[0]].append(p[1])
    for rep, mem in clusters.items():
        mine = [m for m in mem if m.startswith(FOCUS + "_")]
        if not mine:
            continue
        cls = {tid2c[m] for m in mem if m in tid2c}
        c = ("AC" if "AC" in cls else "CC" if "CC" in cls else "unknown")
        for m in mine:
            comp[m] = c
    n_ac = sum(1 for v in comp.values() if v == "AC")
    n_cc = sum(1 for v in comp.values() if v == "CC")
    n_tot = len(comp)
    print(f"genes with a resolved donor compartment: AC {n_ac}, CC {n_cc}, "
          f"unknown {n_tot - n_ac - n_cc}  (of {n_tot})")

    # ---- (a) + (b) ----------------------------------------------------------------
    print(f"\n{'category':30s} {'BMU02483':>16s} {'env pooled':>16s} {'OR':>6s} {'P':>10s}"
          f" {'env median':>11s} {'env range':>15s} | {'AC of them':>11s} {'vs genome':>10s}")
    rows = []
    for name, pred in CATS.items():
        kb = sum(1 for g, d in foc_pfam.items() if pred(d))
        ke = sum(1 for _, p, _ in env for g, d in p.items() if pred(d))
        Ne = sum(n for _, _, n in env)
        orr, p = st.fisher_exact([[kb, foc_n - kb], [ke, Ne - ke]])
        props = [100 * sum(1 for g, d in p.items() if pred(d)) / n for _, p, n in env]
        genes = {g for g, d in foc_pfam.items() if pred(d)}
        ac = sum(1 for g in genes if comp.get(g) == "AC")
        res = sum(1 for g in genes if comp.get(g) in ("AC", "CC"))
        orr2, p2 = st.fisher_exact([[ac, res - ac], [n_ac, n_cc]])
        rows.append((name, kb, foc_n, ke, Ne, orr, p, statistics.median(props),
                     (min(props), max(props)), ac, res, orr2, p2))
        print(f"{name:30s} {kb:6d}/{foc_n:<9d} {ke:8d}/{Ne:<7d} {orr:6.2f} {p:10.2e}"
              f" {statistics.median(props):10.2f}% {min(props):6.2f}-{max(props):5.2f}%"
              f" | {ac:4d}/{res:<6d} {orr2:6.2f} P={p2:.1e}")

    # ---- write the sheet ----------------------------------------------------------
    import openpyxl
    XLSX = "00.manuscript/supplementary/suppl_table.BMU02483_gene_content.xlsx"
    SHEET = "Clinical_function_vs_env"
    wb = openpyxl.load_workbook(XLSX)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)
    ws.append(["Functional_category", "N_BMU02483", "N_genes_BMU02483", "Pct_BMU02483",
               "N_env_pooled", "N_genes_env_pooled", "Pct_env_pooled",
               "OR_BMU02483_vs_env", "Fisher_P", "Env_median_pct", "Env_min_pct", "Env_max_pct",
               "N_AC_associated", "N_compartment_resolved", "Pct_AC_of_resolved",
               "OR_AC_vs_genome", "P_AC_vs_genome"])
    for (name, kb, nb, ke, ne, orr, p, med, rng, ac, res, orr2, p2) in rows:
        ws.append([name, kb, nb, round(100 * kb / nb, 3), ke, ne, round(100 * ke / ne, 3),
                   round(orr, 3), p, round(med, 3), round(rng[0], 3), round(rng[1], 3),
                   ac, res, round(100 * ac / res, 2) if res else None,
                   round(orr2, 3), p2])
    wb.save(XLSX)
    print(f"\n{len(rows)} categories written to [{SHEET}]")


if __name__ == "__main__":
    main()
