#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Build the per-gene sheet for Supplementary Table S14 (BMU02483 gene content).

One row per BMU02483 gene whose donors (>=99% identity over 211 isolates) include clinical
isolates and include none of the 17 environmental isolates closest to BMU02483 in the
core-genome tree order -- the 1,138 genes its own clade lacks. For each gene the row gives
the donor counts, the chromosomal compartment its donors occupy, the accessory chromosomes
involved with their block/mosaic class, and, for core-chromosome donors, how far the donor
copy sits from the nearest chromosome end.

Sheet is appended to 00.manuscript/supplementary/suppl_table.S14_BMU02483_gene_content.xlsx.
"""

import collections
import csv

import openpyxl

TGS = "06.accessory-chrs/tgs_assemblies"
ANN = "02.annotation/per_isolates"
CLASS = "06.accessory-chrs/kmer_analysis/kmer_type.classification.v3.tsv"
XLSX = "00.manuscript/supplementary/suppl_table.S14_BMU02483_gene_content.xlsx"
SHEET = "Clinical_only_gene_routes"
ISO = ["BMU01933", "BMU02028", "BMU02474", "BMU03267", "BMU06673",
       "BMU12683", "BMU12684", "BMU12685", "BMU12686",
       "Fu3.1", "Fu6.1", "SB1", "SX13"]

aclass = {r["AC"]: r["kmer_type"] for r in csv.DictReader(open(CLASS), delimiter="\t")}
gene, clen = {}, {}
for iso in ISO:
    acc = set(open(f"{TGS}/{iso}.accessory_chrs.id").read().split())
    for line in open(f"{ANN}/{iso}/repeatmasker/{iso}.assembly.fasta.masked.length"):
        q = line.split()
        if len(q) >= 2:
            clen[q[0]] = int(q[1])
    for line in open(f"{ANN}/{iso}/funannotate/predict_results/fusarium.gff3"):
        if line.startswith("#"):
            continue
        f = line.rstrip().split("\t")
        if len(f) < 9 or f[2] != "mRNA":
            continue
        for at in f[8].split(";"):
            if at.startswith("ID="):
                tid = at[3:]
                break
        else:
            continue
        L = clen.get(f[0], 0)
        gene[tid] = (f[0], f[0] if f[0] in acc else None, min(int(f[3]), max(L - int(f[4]), 0)))

grp = {r["BMU_ID"]: r for r in csv.DictReader(
    open("05.pan-genome/bmu_ref_groups.211.txt"), delimiter="\t")}
names = sorted([l.strip() for l in open("05.pan-genome/sample.solaniSC.211.list") if l.strip()],
               key=len, reverse=True)


def iso_of(g):
    for n in names:
        if g.startswith(n + "_"):
            return n
    return None


order = [l.strip() for l in open(f"{TGS}/tree_order_allSC.txt") if l.strip()]
b = order.index("BMU02483")
near = order[max(0, b - 10):b + 11]
NEAR = {s for s in near if grp.get(s, {}).get("Species") == "falciformis"
        and grp[s]["MedEnv"] == "env"}
CLIN = {k for k, v in grp.items() if v["MedEnv"] == "med"}

clusters = collections.defaultdict(list)
for line in open("tmp/mmseqs99/out99_cluster.tsv"):
    p = line.rstrip().split("\t")
    if len(p) > 1:
        clusters[p[0]].append(p[1])

rows = []
for rep, mem in clusters.items():
    mine = [m for m in mem if m.startswith("BMU02483_")]
    if not mine:
        continue
    others = [m for m in mem if not m.startswith("BMU02483_")]
    isos = {iso_of(m) for m in others}
    isos.discard(None)
    if (isos & NEAR) or not (isos & CLIN):
        continue
    acm = [m for m in others if m in gene and gene[m][1]]
    ccm = [m for m in others if m in gene and not gene[m][1]]
    comp = ("AC + CC" if acm and ccm else "AC only" if acm
            else "CC only" if ccm else "unresolved (draft donors)")
    acs = sorted({gene[m][1] for m in acm})
    acdesc = ";".join(f"{a}({aclass.get(a,'?').replace('type1_stable','block').replace('type2_mosaic','mosaic')})"
                      for a in acs)
    dcc = min((gene[m][2] for m in ccm), default=None)
    for g in mine:
        rows.append({
            "BMU02483_gene": g,
            "n_donor_isolates": len(isos),
            "n_clinical_donors": len(isos & CLIN),
            "n_near_env_donors": len(isos & NEAR),
            "compartment": comp,
            "AC_donors_class": acdesc,
            "CC_donor_min_dist_to_chromosome_end_bp": dcc,
        })

wb = openpyxl.load_workbook(XLSX)
if SHEET in wb.sheetnames:
    del wb[SHEET]
ws = wb.create_sheet(SHEET)
cols = list(rows[0].keys())
ws.append(cols)
for r in rows:
    ws.append([r[c] for c in cols])
wb.save(XLSX)
print(f"{len(rows)} genes written to {XLSX} [{SHEET}]")
print("compartment counts:", collections.Counter(r["compartment"] for r in rows).most_common())
