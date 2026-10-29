#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rebuild the two HOG sheets of Supplementary Table S14 on the CURRENT definition, so the
table matches the main text (170 BMU02483-specific HOGs, absent from the 17 environmental
isolates closest to BMU02483).

Why this was needed
    The workbook was written from bm2483_specific_hogs_annotated.tsv, which holds only the
    145 HOGs that received an annotation, and its enrichment sheet used 145 as the
    denominator. The analysis itself (bm2483_gene_content_stats.tsv, and the *_near_env
    enrichment file) is on 170 HOGs / 17 isolates, which is what the text reports
    (21/170 = 12.4%, odds ratio 4.3, P = 1.6e-7).

Sheets rewritten
    BMU02483_specific_HOGs  170 rows; annotation columns are '-' for the 25 HOGs that carry
                            none, so the row count matches the text
    Clinical_AC_enrichment  headline test from bm2483_gene_content_stats.tsv, the four
                            presence-threshold tests from the *_near_env file (n = 170),
                            and the all-carried / environmental-carried controls
"""

import csv
import os

import openpyxl

OUT = "06.accessory-chrs/BMU02483_gene_content"
XLSX = "00.manuscript/supplementary/suppl_table.BMU02483_gene_content.xlsx"


def read_tsv(p):
    return list(csv.DictReader(open(p), delimiter="\t"))


specific = read_tsv(f"{OUT}/bm2483_specific_hogs.tsv")
annot = {r["HOG_ID"]: r for r in read_tsv(f"{OUT}/bm2483_specific_hogs_annotated.tsv")}
stats = {r: v for r, v in
         (l.rstrip("\n").split("\t") for l in open(f"{OUT}/bm2483_gene_content_stats.tsv"))}
near = read_tsv(f"{OUT}/bm2483_enrichment_results_near_env.tsv")
old = read_tsv(f"{OUT}/bm2483_enrichment_results.tsv")

print(f"specific HOGs: {len(specific)}   annotated: {len(annot)}   "
      f"stats: env={stats['falc_env_count']} specific={stats['bm2483_specific_hogs']} "
      f"in_clinicalAC={stats['specific_in_clinicalAC']}")

wb = openpyxl.load_workbook(XLSX)

# ---- sheet 1 -------------------------------------------------------------------
del wb["BMU02483_specific_HOGs"]
ws = wb.create_sheet("BMU02483_specific_HOGs", 0)
ws.append(["HOG_ID", "In_clinical_AC_gene_set", "Gene_name", "Pfam_domains"])
for r in sorted(specific, key=lambda x: x["HOG_ID"]):
    hid = r["HOG_ID"]
    a = annot.get(hid, {})
    flag = "Yes" if r["in_clinical_AC"] in ("True", "true", "Yes") else "No"
    ws.append([hid, flag, a.get("Gene_Name", "-") or "-", a.get("Pfam_Domain", "-") or "-"])
print(f"  sheet1 written: {ws.max_row - 1} HOGs")

# ---- sheet 2 -------------------------------------------------------------------
del wb["Clinical_AC_enrichment"]
ws2 = wb.create_sheet("Clinical_AC_enrichment", 1)
ws2.append(["Analysis", "Target_set", "N_target", "N_specific_HOGs",
            "N_specific_in_target", "Odds_ratio", "Fisher_P"])
ws2.append(["specific_vs_clinicalAC_gene_set", "clinical_AC_gene_set",
            int(stats["clinical_ac_hog_set"]), int(stats["bm2483_specific_hogs"]),
            int(stats["specific_in_clinicalAC"]),
            round(float(stats["fisher_specific_odds"]), 3),
            float(stats["fisher_specific_p"])])
for r in near:
    ws2.append([r["analysis"], r["target_set"], int(r["target_n"]), int(r["specific_n"]),
                int(r["specific_in_target"]), round(float(r["odds_ratio"]), 3),
                float(r["fisher_p"])])
for name in ("all_carried_vs_rare_clinicalAC", "env_carried_vs_rare_clinicalAC_control"):
    r = [x for x in old if x["analysis"] == name][0]
    ws2.append([r["analysis"], r["target_set"], int(r["target_n"]), int(r["specific_n"]),
                int(r["specific_in_target"]), round(float(r["odds_ratio"]), 3),
                float(r["fisher_p"])])
print(f"  sheet2 written: {ws2.max_row - 1} tests")

wb.save(XLSX)
print("saved", XLSX)
