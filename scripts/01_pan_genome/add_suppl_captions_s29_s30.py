#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Insert the S29 and S30 supplementary-figure captions into
00.manuscript/supplementary/solani-genome-supplementary-figure-Jun14-GS.docx

Placement: immediately after the S27 caption (the last figure caption) and
before the first table caption, matching the document's existing order.

Every existing caption is a single 'Normal' paragraph with one run in
Times New Roman 11 pt (139700 EMU), so the new ones are built the same way.
Nothing is hand-edited: the script asserts the anchor text, reports what it
inserted, and refuses to run twice.
"""

import os
import shutil
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

DOC = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "00.manuscript", "supplementary",
                   "solani-genome-supplementary-figure-Jun14-GS.docx")

CAPTIONS = [
    (
        "S29",
        "Supplementary Figure S29. Distribution of Starship Captain (DUF3435) and spok "
        "spore-killer genes between accessory and core chromosomes. a, Proportion of each "
        "gene family located on accessory chromosomes (AC), against the expectation implied "
        "by the AC share of the genome (dashed line, 10.8%). Labels give the number of genes "
        "on AC as a fraction of the family total; P values are one-sided binomial tests "
        "against that expectation. b, Counts of Captain and spok genes on core and accessory "
        "chromosomes in each of the 13 chromosome-level assemblies. c, Enrichment of spok "
        "homologues on AC as a function of the protein-identity threshold used to define a "
        "homologue (observed/expected density with 95% confidence intervals), with the "
        "Captain family shown for comparison; thresholds at which no AC-borne spok homologue "
        "was detected are not plotted."
    ),
    (
        "S30",
        "Supplementary Figure S30. Structural context of accessory-chromosome spok "
        "homologues in the FSSC. a, Gene map of every accessory-chromosome contig that "
        "carries a spok homologue; grey bars represent contigs and ticks mark the positions "
        "of spok (green) and Captain/DUF3435 (tan) genes. Only 3 of the 11 contigs also "
        "carry a Captain gene. b, Comparison of the Starship/spok maintenance mechanism "
        "demonstrated in F. oxysporum (left) with the pattern observed here (right): the "
        "spok driver is enriched on accessory chromosomes whereas its Starship marker is "
        "not, indicating that the two are decoupled in the FSSC."
    ),
]


def add_caption(before_par, text):
    """Insert one Normal / Times New Roman 11 pt paragraph before `before_par`."""
    p = before_par.insert_paragraph_before("")
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    return p


def main():
    if not os.path.exists(DOC):
        raise SystemExit(f"docx not found: {DOC}")
    d = docx.Document(DOC)

    # refuse to double-insert
    joined = "\n".join(p.text for p in d.paragraphs)
    already = [tag for tag, _ in CAPTIONS if f"Supplementary Figure {tag}." in joined]
    if already:
        raise SystemExit(f"already present, nothing to do: {already}")

    # anchor = the first paragraph that is a TABLE caption (S7), after the S27 caption
    anchor = next((p for p in d.paragraphs
                   if p.text.strip().startswith("Supplementary Table S7.")), None)
    if anchor is None:
        raise SystemExit("anchor 'Supplementary Table S7.' not found -- refusing to guess")
    s27 = next((p for p in d.paragraphs
                if p.text.strip().startswith("Supplementary Figure S27.")), None)
    if s27 is None:
        raise SystemExit("S27 caption not found -- refusing to guess")

    print(f"anchor  : {anchor.text[:70]!r}")
    print(f"after   : {s27.text[:70]!r}")
    for tag, text in CAPTIONS:
        add_caption(anchor, text)
        print(f"inserted: {tag}  ({len(text)} chars)")

    d.save(DOC)

    # verify by reading back
    d2 = docx.Document(DOC)
    ok = 0
    for i, p in enumerate(d2.paragraphs):
        t = p.text.strip()
        for tag, text in CAPTIONS:
            if t.startswith(f"Supplementary Figure {tag}."):
                r = p.runs[0]
                print(f"  verified par[{i}] {tag}: style={p.style.name!r} "
                      f"font={r.font.name!r} size={r.font.size.pt if r.font.size else None}pt "
                      f"len={len(t)}")
                ok += 1
    print(f"\n{ok}/{len(CAPTIONS)} captions verified in the saved file")
    return 0 if ok == len(CAPTIONS) else 1


if __name__ == "__main__":
    sys.exit(main())
