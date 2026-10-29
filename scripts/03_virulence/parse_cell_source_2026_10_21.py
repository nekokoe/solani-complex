#!/usr/bin/env python3
"""Canonical parser for the 2026-10-21 cell source data.

Supersedes `parse_cell_assay.py` (which read the 2026-10-12 workbook): from 2026-10-22
the whole cell arm — Supplementary Fig. S28 and Supplementary Table S12 — is derived
from `261020细胞蜡螟结果_and_like方法修改意见/Source data-cell and larvae.xlsx`.

Layout of sheet `cell` (verified cell by cell, see verify_cell_source_data.py)
    block header row   PBMCs-derived macrophage / THP1-derived macrophage
    cytokine row       TNFa / IL1b / csf2 / cxcl1
    assay row          qPCR / ELISA  + groups: Ctrl, LPS, then the 21 strains
    `High concentration` row, `Low concentration` row
    Each group spans THREE ADJACENT COLUMNS = the three experimental replicates.
    Only IL-1b and TNF-a have an ELISA block.

Two things this parser has to get right
  * replicates are columns here (in the 2026-10-12 workbook they were rows);
  * the low-concentration rows carry no Ctrl/LPS. The control pair is SHARED between
    the two concentrations — `parse_cell_assay.py` verified this on the old workbook,
    where the high and low Ctrl/LPS values are byte-identical — so the high-block Ctrl
    is carried over to the low block. Without that, fold-change would be undefined at
    the low dose.

Outputs (same schema as before, plus `model`, so downstream readers keep working)
    07.lethal_groups/cell_assay/cell_assay_long.tsv
    07.lethal_groups/cell_assay/cell_assay_summary.tsv
"""

import os

import numpy as np
import openpyxl
import pandas as pd

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "261020细胞蜡螟结果_and_like方法修改意见",
                    "Source data-cell and larvae.xlsx")
OUT = os.path.join(ROOT, "07.lethal_groups/cell_assay")
CONC = {"High concentration": 1e5, "Low concentration": 1e3}
ASSAY = {"qPCR": "mRNA", "ELISA": "protein"}


def parse_sheet():
    ws = openpyxl.load_workbook(XLSX, data_only=True)["cell"]
    model = cytokine = None
    header, rows = {}, []
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a is None:
            continue
        a = str(a).strip()
        if a in ("PBMCs-derived macrophage", "THP1-derived macrophage"):
            model = "PBMC" if a.startswith("PBMC") else "THP1"
        elif a in ("TNFa", "IL1b", "csf2", "cxcl1"):
            cytokine = a
        elif a in ASSAY:
            assay = ASSAY[a]
            header = {}
            for c in range(2, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v not in (None, ""):
                    header[c] = str(v).strip()
        elif a in CONC:
            conc = CONC[a]
            for c, name in header.items():
                vals = [ws.cell(r, c + k).value for k in range(3)]
                vals = [float(v) for v in vals if isinstance(v, (int, float))]
                for v in vals:
                    rows.append(dict(model=model, cytokine=cytokine, assay=assay,
                                     concentration=conc, unit=name, value=v))
    return pd.DataFrame(rows)


def main():
    long = parse_sheet()
    long["is_control"] = long.unit.isin(["Ctrl", "LPS"])
    long["strain"] = np.where(long.is_control, pd.NA, long.unit)
    os.makedirs(OUT, exist_ok=True)
    long.to_csv(f"{OUT}/cell_assay_long.tsv", sep="\t", index=False)

    print(f"rows: {len(long)}")
    print(f"models: {sorted(long.model.unique())}")
    print(f"cytokines: {sorted(long.cytokine.unique())}")
    print(f"ELISA available for: {sorted(long[long.assay == 'protein'].cytokine.unique())}")
    print(f"replicates per strain x assay x conc: "
          f"{sorted(long[~long.is_control].groupby(['model','cytokine','assay','concentration','strain']).size().unique())}")

    # the low block has no Ctrl/LPS -> carry the high-block control pair over
    ctrl = long[long.is_control]
    have = set(zip(ctrl.model, ctrl.cytokine, ctrl.assay, ctrl.concentration))
    need = set(zip(long.model, long.cytokine, long.assay, long.concentration))
    carry = sorted({(m, c, a) for (m, c, a, k) in need - have})
    print(f"\nblocks without their own control (carried from 10^5): {len(carry)}")
    for m, c, a in carry:
        print(f"   {m:5s} {c:6s} {a:7s}")
    if carry:
        add = ctrl[(ctrl.concentration == 1e5) &
                   ctrl.set_index(["model", "cytokine", "assay"]).index.isin(carry)].copy()
        add["concentration"] = 1e3
        long = pd.concat([long, add], ignore_index=True)

    base = (long[long.is_control & (long.unit == "Ctrl")]
            .groupby(["model", "cytokine", "assay", "concentration"]).value.mean()
            .rename("ctrl_mean"))
    s = (long[~long.is_control]
         .groupby(["model", "cytokine", "assay", "concentration", "strain"]).value
         .agg(["mean", "std", "size"]).reset_index()
         .merge(base.reset_index(), on=["model", "cytokine", "assay", "concentration"]))
    s["fold_vs_ctrl"] = s["mean"] / s.ctrl_mean
    s.to_csv(f"{OUT}/cell_assay_summary.tsv", sep="\t", index=False)
    print(f"\nsummary rows: {len(s)}")

    hi = s[s.concentration == 1e5]
    print("\nfold change vs Ctrl at 10^5, by model x cytokine x assay:")
    print(hi.groupby(["model", "cytokine", "assay"]).fold_vs_ctrl
          .describe()[["count", "mean", "min", "max"]].round(2).to_string())
    print(f"\nwrote {OUT}/cell_assay_long.tsv and cell_assay_summary.tsv")


if __name__ == "__main__":
    main()
